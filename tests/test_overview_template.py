"""Market Overview template: snapshot, Excel export fidelity, collapsible page, and navigation."""

from __future__ import annotations

import io
import re
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from market_intelligence import overview_snapshot as snapshot_module
from market_intelligence.cross_asset_universe import CURRENCY_VS_USD, FX_PAIRS, INSTRUMENT_BY_ID, MOVE_INSTRUMENT, YAHOO_CROSS_ASSET
from market_intelligence.ibkr_live_universe import STOCK_GROUPS, SUBSECTOR_ETFS, subsector_groups
from market_intelligence.ingest_yahoo_cross_asset import incremental_start
from market_intelligence.overview_export import (
    EXCEL_MIME,
    METADATA_COLUMN,
    PERCENT_FORMAT,
    TEMPLATE_PATH,
    TEMPLATE_ROWS,
    TEMPLATE_SHEET,
    export_filename,
    fill_template,
    level_number_format,
)
from market_intelligence.overview_metrics import percent_change_of_level, risk_window_stats
from market_intelligence.overview_snapshot import (
    SECTION_CREDIT,
    SECTION_FOREX,
    SECTION_ORDER,
    SECTION_SECTORS,
    SECTION_US_INDEXES,
    SECTION_YIELD_CURVE,
    SUBSECTOR_GROUP_BY_SECTOR,
    overview_snapshot,
    section_by_id,
)
from market_intelligence.page_registry import NAV_SECTIONS, PAGE_BY_ROUTE, RETIRED_ROUTE_IDS, specs_by_section
from market_intelligence.sector_mapping import CANONICAL_SECTORS
from market_intelligence.taxonomy import SECTOR_PROXIES

ROOT = Path(__file__).resolve().parents[1]
TODAY = date(2026, 10, 6)


class _FakeConn:
    @contextmanager
    def begin_nested(self):
        yield self


def _daily(start: date, count: int, base: float, step: float) -> list[date]:
    """``count`` consecutive calendar days ending on ``TODAY`` when ``start`` is TODAY - (count - 1)."""
    return [start + timedelta(days=offset) for offset in range(count)]


def _monitor_history(symbols: list[str]) -> dict:
    days = _daily(TODAY - timedelta(days=400), 401, 100.0, 0.1)
    bars: dict = {}
    returns: dict = {}
    latest: dict = {}
    meta: dict = {}
    for index, symbol in enumerate(symbols):
        series = [(day, 100.0 + index + 0.05 * offset) for offset, day in enumerate(days)]
        bars[symbol] = [{"date": day.isoformat(), "value": value} for day, value in series]
        latest[symbol] = series[-1][1]
        meta[symbol] = {"latest": series[-1][0].isoformat()}
        returns[symbol] = {"1D": 0.001 * (index + 1), "1W": 0.005, "1M": 0.02, "3M": 0.05, "6M": None, "1Y": None}
    return {"bars": bars, "returns": returns, "latest_price": latest, "meta": meta}


def _eod_closes(symbols: list[str]) -> dict:
    days = _daily(TODAY - timedelta(days=120), 121, 50.0, 0.1)
    out: dict = {}
    for index, symbol in enumerate(symbols):
        prices = {day: 50.0 + index + 0.02 * offset for offset, day in enumerate(days)}
        if symbol == "XLU":
            # One sector without stored history: must stay missing, never zero.
            out[symbol] = {"bars": [], "prices": {}, "adjustment_basis": None, "provider": None, "rejection": "missing"}
            continue
        out[symbol] = {
            "bars": [{"date": day.isoformat(), "value": value} for day, value in sorted(prices.items())],
            "prices": prices,
            "adjustment_basis": "adj_close",
            "provider": "IBKR",
            "rejection": None,
        }
    return out


def _bars() -> list[dict]:
    rows: list[dict] = []
    days = _daily(TODAY - timedelta(days=380), 381, 1.0, 0.0)
    for index, instrument in enumerate(YAHOO_CROSS_ASSET):
        if instrument.instrument_id == "USDCHF":
            continue  # unavailable instrument stays missing
        if instrument.instrument_id == "EURUSD":
            series_days = days[:-10]  # stale: last observation 10 days ago
        else:
            series_days = days
        for offset, day in enumerate(series_days):
            rows.append({"instrument_id": instrument.instrument_id, "bar_date": day, "close": 10.0 + index + 0.01 * offset, "retrieved_at": None})
    return rows


def _observations(series_ids: list[str]) -> dict:
    """FRED series end yesterday; Treasury (UST_*) series carry today's print, as on the host."""
    days = _daily(TODAY - timedelta(days=400), 401, 4.0, 0.0)
    out: dict = {}
    for series_id in series_ids:
        rows = [{"observation_date": day, "value": 4.0 + 0.001 * offset} for offset, day in enumerate(days)]
        if series_id.startswith("DGS"):
            rows = rows[:-1]
        out[series_id] = rows
    return out


def _credit() -> dict:
    return {
        "buckets": [
            {"series_id": "BAMLC0A0CM", "oas_bps": 90.0, "as_of": TODAY.isoformat(), "change_1d_bps": -1.0, "change_1w_bps": 1.0, "change_1m_bps": 2.0, "change_3m_bps": 3.0},
            {"series_id": "BAMLH0A0HYM2", "oas_bps": 310.0, "as_of": TODAY.isoformat(), "change_1d_bps": 2.0, "change_1w_bps": 3.0, "change_1m_bps": 4.0, "change_3m_bps": 5.0},
        ]
    }


def _vol(metric_ids: list[str]) -> dict:
    days = _daily(TODAY - timedelta(days=100), 101, 15.0, 0.0)
    return {metric: [{"as_of": day, "value": 15.0 + index + 0.01 * offset} for offset, day in enumerate(days)] for index, metric in enumerate(metric_ids)}


@pytest.fixture
def snapshot(monkeypatch):
    monkeypatch.setattr(snapshot_module, "load_monitor_history", lambda conn, symbols, since=None: _monitor_history(list(symbols)))
    monkeypatch.setattr(snapshot_module, "load_equity_eod_closes", lambda conn, symbols, since=None: _eod_closes(list(symbols)))
    monkeypatch.setattr(snapshot_module.read_models, "credit_context", lambda conn: _credit())
    monkeypatch.setattr(snapshot_module.read_models, "recent_observations", lambda conn, ids, limit=40: _observations(list(ids)))
    monkeypatch.setattr(snapshot_module.read_models, "metric_history", lambda conn, metric_id, limit=40: [{"as_of": TODAY - timedelta(days=offset), "value": 30.0 - offset * 0.1} for offset in range(200)])
    monkeypatch.setattr(snapshot_module.read_models, "yahoo_vol_metric_history", lambda conn, ids, since=None: _vol(list(ids)))
    monkeypatch.setattr(snapshot_module, "yahoo_cross_asset_bars", lambda conn, since=None: _bars())
    return overview_snapshot(_FakeConn(), today=TODAY)


# ---- instruments and registry ----------------------------------------------------------


def test_usdcnh_and_move_are_registered_with_correct_orientation():
    cnh = INSTRUMENT_BY_ID["USDCNH"]
    assert cnh.yahoo_symbol == "CNH=F"
    assert cnh.base_currency == "USD" and cnh.quote_currency == "CNH"
    assert cnh.invert_vs_usd is True, "USD/CNH rising = USD appreciation; the vs-USD view inverts it like USD/JPY"
    assert ("USDCNH", "CNH") in CURRENCY_VS_USD
    assert ("USDCNH", "USD/CNH") in FX_PAIRS
    assert MOVE_INSTRUMENT.yahoo_symbol == "^MOVE" and MOVE_INSTRUMENT.instrument_id == "MOVE"
    assert MOVE_INSTRUMENT in YAHOO_CROSS_ASSET
    assert "CNY" not in cnh.display_name


def test_incremental_refresh_requests_provider_max_for_unseen_instruments():
    assert incremental_start(None, today=TODAY) is None
    assert incremental_start(TODAY - timedelta(days=1), today=TODAY) < TODAY - timedelta(days=1)


def test_sidebar_registry_regroups_markets_and_retires_entries():
    assert NAV_SECTIONS == ("Markets", "Positioning", "Economy", "Research", "System")
    grouped = specs_by_section()
    markets = [spec.title for spec in grouped["Markets"]]
    assert markets[0] == "Market Overview"
    assert {"FOREX", "Commodities", "Crypto", "US Markets", "Global Markets", "Rates & Curve", "Credit", "Options & Volatility"} <= set(markets)
    visible_titles = {spec.title for specs in grouped.values() for spec in specs}
    for retired in ("Equities & Sectors", "Bond Trading Activity", "Bond Research", "Power Producers", "Morning Brief", "Methodology"):
        assert retired not in visible_titles
    assert set(RETIRED_ROUTE_IDS) == {"sectors", "order_flow", "fixed_income", "power_producers", "morning_brief", "methodology"}
    for route_id in RETIRED_ROUTE_IDS:
        assert PAGE_BY_ROUTE[route_id].hidden and PAGE_BY_ROUTE[route_id].file_path
    assert PAGE_BY_ROUTE["overview"].default
    assert len({spec.url_path for spec in PAGE_BY_ROUTE.values()}) == len(PAGE_BY_ROUTE)


def test_sector_drilldowns_map_every_canonical_sector_to_a_subsector_group():
    names = {name for name, _members in subsector_groups()}
    for sector in CANONICAL_SECTORS:
        assert SUBSECTOR_GROUP_BY_SECTOR[sector] in names, sector
    assert SUBSECTOR_GROUP_BY_SECTOR["Technology"] == "Tech"
    power = [name for name, _members in STOCK_GROUPS if name.startswith("Power")]
    assert len(power) >= 4, "Power instrument groups must survive the sidebar cleanup"
    assert ("GRID", "grid/transmission infrastructure", "Utilities") in SUBSECTOR_ETFS


# ---- metrics and snapshot ---------------------------------------------------------------


def test_risk_window_stats_require_full_window_and_never_zero_fill():
    days = _daily(TODAY - timedelta(days=100), 64, 0, 0)
    points = [(day, 100.0 * (1.01 ** offset)) for offset, day in enumerate(days)]
    stats = risk_window_stats(points)
    assert stats["max_dd"] == pytest.approx(0.0)
    assert stats["vol_ann"] == pytest.approx(0.0, abs=1e-9)
    short = risk_window_stats(points[:30])
    assert short == {"vol_ann": None, "sharpe": None, "max_dd": None}
    assert percent_change_of_level(4.0, 4.0) == pytest.approx(0.04 / 3.96)
    assert percent_change_of_level(None, 1.0) is None


def test_snapshot_matches_template_sections_rows_and_statuses(snapshot):
    assert [section["section_id"] for section in snapshot["sections"]] == list(SECTION_ORDER)
    for section_id, rows in TEMPLATE_ROWS.items():
        section = section_by_id(snapshot, section_id)
        assert [row["key"] for row in section["rows"]] == list(rows)
    forex = section_by_id(snapshot, SECTION_FOREX)
    by_key = {row["key"]: row for row in forex["rows"]}
    assert by_key["USDCNH"]["status"] == "ok" and by_key["USDCNH"]["level"] is not None
    assert by_key["USDCHF"]["status"] == "missing" and by_key["USDCHF"]["level"] is None
    assert by_key["EURUSD"]["status"] == "stale"
    yields = section_by_id(snapshot, SECTION_YIELD_CURVE)
    move = next(row for row in yields["rows"] if row["key"] == "MOVE")
    assert move["status"] == "ok" and move["drill"]["anchor"] == "move-index"
    ten = next(row for row in yields["rows"] if row["key"] == "10Y")
    assert ten["as_of"] == TODAY.isoformat() and ten["source"] == "TREASURY", "Treasury print wins over the lagging FRED copy"
    assert ten["level"] == pytest.approx(4.4)
    assert ten["changes"]["1D"] == pytest.approx(0.1) and ten["changes_pct"]["1D"] == pytest.approx(0.001 / 4.399)
    assert ten["changes"]["1W"] == pytest.approx(0.7) and ten["changes"]["1Y"] == pytest.approx(36.5)
    assert ten["changes"]["6M"] is not None
    slope = next(row for row in yields["rows"] if row["key"] == "2s10s")
    # all synthetic tenors share one series, so 10Y - 2Y = 0 bps on the common (Treasury) date and the stored metric history is not used
    assert slope["level"] == pytest.approx(0.0) and slope["as_of"] == TODAY.isoformat() and slope["changes"]["1W"] == pytest.approx(0.0)
    fly = next(row for row in yields["rows"] if row["key"] == "2s5s10s")
    assert fly["level"] == pytest.approx(0.0) and fly["status"] == "ok"
    sectors = section_by_id(snapshot, SECTION_SECTORS)
    utilities = next(row for row in sectors["rows"] if row["label"] == "Utilities")
    assert utilities["status"] == "missing" and utilities["changes"]["1D"] is None
    tech = next(row for row in sectors["rows"] if row["label"] == "Technology")
    assert tech["drill"] == {"route_id": "us_markets", "anchor": "subsector-performance", "state": {"us_subsector_sector": "Tech"}, "sector": "Technology"}
    assert tech["key"] == SECTOR_PROXIES["Technology"]
    assert snapshot["rows_missing"] >= 2 and snapshot["rows_stale"] >= 1
    assert snapshot["read_errors"] == {}
    assert re.fullmatch(r"[0-9a-f]{16}", snapshot["snapshot_id"])


def test_snapshot_isolates_a_failing_read(monkeypatch, snapshot):
    def boom(conn):
        raise RuntimeError("relation missing")

    monkeypatch.setattr(snapshot_module.read_models, "credit_context", boom)
    degraded = overview_snapshot(_FakeConn(), today=TODAY)
    assert degraded["read_errors"] == {"credit": "RuntimeError"}
    credit = section_by_id(degraded, SECTION_CREDIT)
    assert all(row["status"] == "missing" for row in credit["rows"])
    assert section_by_id(degraded, SECTION_US_INDEXES)["rows"][0]["status"] == "ok"


# ---- Excel export --------------------------------------------------------------------------


def test_template_asset_is_packaged_and_unmodified_in_shape():
    assert TEMPLATE_PATH.exists() and TEMPLATE_PATH.is_relative_to(ROOT)
    book = load_workbook(TEMPLATE_PATH)
    assert book.sheetnames == [TEMPLATE_SHEET]
    sheet = book[TEMPLATE_SHEET]
    assert sheet.max_row == 98 and sheet.max_column == 9
    assert not sheet.merged_cells.ranges
    assert sheet["A1"].value == "US Indexes" and sheet["A94"].value.startswith("USD/CNH") and sheet["A45"].value.startswith("MOVE")


def test_export_fills_every_field_and_preserves_template_layout(snapshot):
    exported_at = datetime(2026, 10, 6, 17, 5, tzinfo=ZoneInfo("America/New_York"))
    payload = fill_template(snapshot, exported_at=exported_at)
    template = load_workbook(TEMPLATE_PATH)[TEMPLATE_SHEET]
    book = load_workbook(io.BytesIO(payload))
    assert book.sheetnames == [TEMPLATE_SHEET]
    sheet = book[TEMPLATE_SHEET]
    assert not sheet.merged_cells.ranges
    for letter, dimension in template.column_dimensions.items():
        assert sheet.column_dimensions[letter].width == dimension.width
    for row_index, dimension in template.row_dimensions.items():
        assert sheet.row_dimensions[row_index].height == dimension.height
    # headers and description rows untouched
    for excel_row in (1, 3, 13, 19, 32, 47, 52, 58, 60, 76, 87, 96):
        for column in "ABCDEFGHI":
            assert sheet["{0}{1}".format(column, excel_row)].value == template["{0}{1}".format(column, excel_row)].value
    # SPY: numeric level with label in format, real date, fractional returns formatted as percent
    assert isinstance(sheet["A2"].value, (int, float)) and sheet["A2"].number_format == level_number_format("SPY", "price")
    assert isinstance(sheet["B2"].value, datetime) and sheet["B2"].value.date() == TODAY
    assert sheet["C2"].value == pytest.approx(0.001) and sheet["C2"].number_format == PERCENT_FORMAT
    assert isinstance(sheet["G2"].value, (int, float)) and isinstance(sheet["I2"].value, (int, float))
    # Yield curve: %Change columns carry the fractional change of the level, 6 horizons
    ten = TEMPLATE_ROWS[SECTION_YIELD_CURVE]["10Y"]
    assert sheet["A{0}".format(ten)].value == pytest.approx(4.4) and sheet["A{0}".format(ten)].number_format.endswith('0.00"%"')
    assert sheet["C{0}".format(ten)].value == pytest.approx(0.001 / 4.399)
    assert sheet["H{0}".format(ten)].value is not None
    # Credit: basis-point moves, numeric
    ig = TEMPLATE_ROWS[SECTION_CREDIT]["BAMLC0A0CM"]
    assert sheet["A{0}".format(ig)].value == 90.0 and sheet["C{0}".format(ig)].value == -1.0
    em = TEMPLATE_ROWS[SECTION_CREDIT]["BAMLEMCBPIOAS"]
    assert sheet["A{0}".format(em)].value == "EM OAS - n/a" and sheet["B{0}".format(em)].value is None and sheet["C{0}".format(em)].value is None
    # USD/CNH and MOVE populated; USD/CHF missing stays blank (never zero)
    assert isinstance(sheet["A94"].value, (int, float)) and sheet["A94"].number_format.startswith('"USD/CNH - "')
    assert isinstance(sheet["A45"].value, (int, float)) and sheet["C45"].value is not None
    assert sheet["A45"].number_format == '"MOVE - "0.00', "MOVE is index points, not a yield"
    assert sheet["A93"].value == "USD/CHF - n/a" and sheet["C93"].value is None
    # Utilities (missing EOD history) stays blank
    assert sheet["A30"].value == "Utilities - n/a" and sheet["C30"].value is None
    # metadata outside A:I
    labels = [sheet["{0}{1}".format(METADATA_COLUMN, index)].value for index in range(1, 20)]
    assert "Exported (America/New_York)" in labels and "Snapshot id" in labels
    assert sheet["L2"].value.startswith("2026-10-06 17:05")
    assert sheet["L4"].value == snapshot["snapshot_id"]
    for row in sheet.iter_rows(min_col=1, max_col=9):
        for cell in row:
            assert not (isinstance(cell.value, str) and cell.value.startswith("=")), "no formulas expected"


def test_export_filename_and_mime():
    stamp = datetime(2026, 10, 6, 21, 5, tzinfo=ZoneInfo("UTC"))
    assert export_filename(stamp) == "Market_Overview_2026-10-06_1705_ET.xlsx"
    assert re.fullmatch(r"Market_Overview_\d{4}-\d{2}-\d{2}_\d{4}_ET\.xlsx", export_filename())
    assert EXCEL_MIME.endswith("spreadsheetml.sheet")


# ---- page --------------------------------------------------------------------------------------


def _overview_app(monkeypatch, snapshot):
    def fake_cached(fn_name, *args, **kwargs):
        if fn_name == "overview_snapshot":
            return snapshot
        raise RuntimeError("unexpected read {0}".format(fn_name))

    monkeypatch.setattr("market_intelligence.ui.cached_read", fake_cached)
    at = AppTest.from_file(str(ROOT / "pages" / "10_Market_Pulse.py"), default_timeout=30)
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    return at


def test_overview_page_renders_template_sections_from_one_snapshot(monkeypatch, snapshot):
    at = _overview_app(monkeypatch, snapshot)
    assert "Market Overview" in [t.value for t in at.title]
    labels = [button.label for button in at.button]
    for title in ("US Indexes", "Market Ratios", "Sectors", "Yield Curve", "Credit Spreads", "VIX Term Structure", "Global Equity Performance", "Commodities", "FOREX", "Crypto"):
        assert title in labels
    assert len(at.dataframe) == len(SECTION_ORDER)
    heads = [h.value for h in at.subheader]
    assert "What matters" not in heads and "Category snapshot" not in heads
    assert any("Export to Excel" in str(getattr(widget, "label", "")) for widget in at.get("download_button")) or "Export to Excel" in str(at)


def test_overview_sections_collapse_independently_and_persist(monkeypatch, snapshot):
    at = _overview_app(monkeypatch, snapshot)
    toggle = next(button for button in at.button if button.key == "overview_toggle_sectors")
    toggle.click().run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.session_state["overview_sections_collapsed"] == {"sectors": True}
    assert len(at.dataframe) == len(SECTION_ORDER) - 1
    at.run()  # plain rerun keeps the collapsed state
    assert len(at.dataframe) == len(SECTION_ORDER) - 1
    other = next(button for button in at.button if button.key == "overview_toggle_forex")
    other.click().run()
    assert at.session_state["overview_sections_collapsed"] == {"sectors": True, "forex": True}
    assert len(at.dataframe) == len(SECTION_ORDER) - 2


def test_sector_tile_primes_subsector_selector_and_switches_page(monkeypatch, snapshot):
    at = _overview_app(monkeypatch, snapshot)
    tiles = [button for button in at.button if button.key and button.key.startswith("overview_sector_tile_")]
    assert [button.label for button in tiles] == list(CANONICAL_SECTORS)
    calls: list = []
    monkeypatch.setattr("market_intelligence.navigation_links.st.switch_page", lambda target: calls.append(target))
    next(button for button in tiles if button.label == "Technology").click().run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.session_state["us_subsector_sector"] == "Tech"
    assert at.session_state["mi_pending_scroll_anchor"] == "subsector-performance"
    assert calls == ["pages/22_US_Markets.py"]


def test_overview_page_does_not_fetch_or_write(monkeypatch, snapshot):
    source = (ROOT / "market_intelligence" / "overview_ui.py").read_text(encoding="utf-8") + (ROOT / "market_intelligence" / "overview_export.py").read_text(encoding="utf-8")
    for forbidden in ("yfinance", "requests.", "ingest_", "INSERT", "UPDATE ", "psycopg", "sqlalchemy"):
        assert forbidden not in source
